#==============================================================================
# 06_exhibits.py
# Structural inequalities in education in China, CFPS 2010 to 2018
#
# Builds the tables and the figure of the paper from the checked output:
# Exhibits/Main tables and figure.docx   Tables I to IV and Figure 1
# Exhibits/Figure 1.png, Figure 1.pdf    Figure 1 in black and white
# Exhibits/Supplementary tables.docx     Tables S1 to S11
#
# The script reads only the folder Output (checked), which holds aggregate
# tables, so it needs neither the CFPS data nor the encrypted disk. Tables
# follow the earlier manuscript (bold table number, italic title, notes that
# start with a bold "Note:") and the rules of Equality, Diversity and
# Inclusion (Roman numerals for tables, three horizontal lines only, the
# figure in black and white).
#
# How to run. In Terminal, from the project folder:
#   python3 "R scripts/06_exhibits.py"
# It needs Python 3 with pandas, matplotlib, lxml and python-docx.
#
# Status. Written on 24 October 2026.
#==============================================================================

import os
import re
from xml.sax.saxutils import escape

import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "Output (checked)")
OUT = os.path.join(ROOT, "Exhibits")

DASH = "–"          # en dash, shown for a value that is not reported
TEXT_WIDTH = 9360        # US Letter with 1-inch margins, in twentieths of a point
INDENT = 227             # 0.4 cm
RULE = '<w:{0} w:val="single" w:sz="18" w:space="0" w:color="000000"/>'
HIDDEN = ("<10", "suppressed")

GROUPS3 = [("agricultural", "Agricultural hukou at age 3"),
           ("non-agricultural", "Non-agricultural hukou at age 3"),
           ("difference", "Difference")]
GROUPS4 = [("all", "All")] + GROUPS3

PROVINCES = {
    "11": "Beijing", "12": "Tianjin", "13": "Hebei", "14": "Shanxi", "15": "Inner Mongolia",
    "21": "Liaoning", "22": "Jilin", "23": "Heilongjiang", "31": "Shanghai", "32": "Jiangsu",
    "33": "Zhejiang", "34": "Anhui", "35": "Fujian", "36": "Jiangxi", "37": "Shandong",
    "41": "Henan", "42": "Hubei", "43": "Hunan", "44": "Guangdong", "45": "Guangxi",
    "46": "Hainan", "50": "Chongqing", "51": "Sichuan", "52": "Guizhou", "53": "Yunnan",
    "54": "Tibet", "61": "Shaanxi", "62": "Gansu", "63": "Qinghai", "64": "Ningxia",
    "65": "Xinjiang"}

# Levels of education (CFPS codes 1 to 5, and 1 to 6 in check R13)
EDU_LEVELS = ["Illiterate or semi-literate", "Primary school", "Junior high school",
              "Senior high school", "College or above"]
EDU6_LEVELS = EDU_LEVELS[:4] + ["Three-year college", "Bachelor's degree or above"]

# Labels of the targets of the average marginal effects, by outcome. A
# target that is a probability is shown in percentage points, and an
# expected value in points of the outcome.
TARGETS = {
    "edu5": {"ge2": "Primary school or above", "ge3": "Junior high school or above",
             "ge4": "Senior high school or above", "ge5": "College or above",
             "mean": "Expected education level, 1 to 5"},
    "edu6": {"ge2": "Primary school or above", "ge3": "Junior high school or above",
             "ge4": "Senior high school or above", "ge5": "Three-year college or above",
             "ge6": "Bachelor's degree or above", "mean": "Expected education level, 1 to 6"},
    "rating": {"ge2": "Rating 2 or above", "ge3": "Rating 3 or above", "ge4": "Rating 4 or above",
               "ge5": "Rating 5 or above", "ge6": "Rating 6 or above", "ge7": "Rating 7",
               "mean": "Expected rating, 1 to 7"},
    "rating0": {"ge1": "Rating 1 or above, that is, not code 79", "ge2": "Rating 2 or above",
                "ge3": "Rating 3 or above", "ge4": "Rating 4 or above", "ge5": "Rating 5 or above",
                "ge6": "Rating 6 or above", "ge7": "Rating 7",
                "mean": "Expected rating, 0 to 7, with code 79 as 0"},
    "lang": {"pr_mandarin": "Mandarin at home", "pr_minority": "Minority language at home",
             "pr_dialect": "Chinese dialect at home"},
    "wordtest": {"mean": "Expected word test score, points"},
    "t1": {"pr": "T1, primary school among all"},
    "t2": {"pr": "T2, junior high school among those with primary school"},
    "t3": {"pr": "T3, senior high school among those with junior high school"},
    "t4": {"pr": "T4, college among those with senior high school"},
    "s1": {"pr": "S1, attended primary school among all"},
    "s2": {"pr": "S2, attended junior high school among those who attended primary school"},
    "s3": {"pr": "S3, attended senior high school among those who attended junior high school"},
    "s4": {"pr": "S4, attended college among those who attended senior high school"},
    "y_sh": {"pr": "Senior high school or above by 2018"},
    "y_ter": {"pr": "College by 2018, completed or enrolled"},
    "enrol": {"pr": "Enrolled in school"},
}

# Quantities of the decomposition by the method of Karlson, Holm and Breen
KHB = {"total": "Total", "direct": "Direct", "indirect": "Indirect, through the measures",
       "pct_indirect": "Indirect as a percentage of the total",
       "contrib_lang_min": "Through minority language at home",
       "contrib_lang_dia": "Through Chinese dialect at home",
       "contrib_language": "Through home language, both categories",
       "contrib_rating": "Through the interviewer's rating of Mandarin",
       "contrib_wordtest": "Through the word test score"}

OUTCOMES = {
    "edu5": "five education levels, ordered logit",
    "edu6": "six education levels, ordered logit",
    "t1": "transition T1, logit", "t2": "transition T2, logit",
    "t3": "transition T3, logit", "t4": "transition T4, logit",
    "s1": "stage S1, logit", "s2": "stage S2, logit", "s3": "stage S3, logit", "s4": "stage S4, logit",
    "lang": "home language, multinomial logit with Chinese dialect as the base",
    "rating": "interviewer's rating of Mandarin, ordered logit",
    "rating0": "interviewer's rating of Mandarin with code 79 as the lowest category, ordered logit",
    "wordtest": "word test score, linear regression",
    "y_sh": "senior high school or above by 2018, logit",
    "y_ter": "college by 2018, logit",
    "enrol": "enrolled in school, logit",
}

TERMS = {"nonhan": "non-Han ethnicity", "agri3": "hukou at age 3", "nh_ag": "the interaction",
         "male": "gender", "paredu": "parents' education", "province": "province",
         "cohort": "birth cohort"}

# Terms in the table of the selection of check R10
R10_TERMS = {"nonhan": "Non-Han ethnicity", "agri3": "Agricultural hukou at age 3",
             "nh_ag": "Non-Han × agricultural hukou at age 3", "male": "Gender",
             "paredu": "Parents' education", "province": "Province", "cohort": "Birth cohort"}

AME_FILES = ("T2_ame_main", "T3_ame_language", "T4_bcohort", "S4_transitions", "S5_ame_other",
             "S6_ame_enrol", "S7_robust")


#------------------------------------------------------------------------------
# Numbers and words
#------------------------------------------------------------------------------
def rd(name):
    return pd.read_csv(os.path.join(SRC, name + ".csv"), encoding="utf-8-sig", dtype=str,
                       keep_default_na=False)


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def fmt(x, d, scale=1.0):
    v = num(x)
    if v is None:
        return DASH
    s = f"{v * scale:.{d}f}"
    if float(s) == 0:
        s = f"{0:.{d}f}"
    return s


def fmt_p(x):
    v = num(x)
    if v is None:
        return DASH
    return "<0.001" if v < 0.001 else f"{v:.3f}"


def fmt_n(x):
    if x == "":
        return DASH
    v = num(x)
    return f"{int(round(v)):,}" if v is not None else x


def cap(s):
    return s[0].upper() + s[1:] if s else s


def join_and(items):
    """a, b and c"""
    items = list(items)
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def is_prob(outcome, target):
    return target not in KHB and target != "mean"


def ame_lines(r):
    """The estimate with its p-value, and the 95% confidence interval, in
    percentage points for a probability and to two decimals for a mean"""
    if r is None or num(r["estimate"]) is None:
        return [DASH]
    prob = is_prob(r["outcome"], r["target"])
    s, d = (100.0, 1) if prob else (1.0, 2)
    return [f"{fmt(r['estimate'], d, s)} ({fmt_p(r['p'])})",
            f"[{fmt(r['ci_low'], d, s)}, {fmt(r['ci_high'], d, s)}]"]


def n_text(r):
    return DASH if r is None else f"{fmt_n(r['n'])} ({fmt_n(r['n_nonhan'])})"


def ame_get(d, **kw):
    m = pd.Series(True, index=d.index)
    for k, v in kw.items():
        m &= d[k] == v
    r = d[m]
    return None if r.empty else r.iloc[0]


def counts_of(d, r):
    """n and non-Han n of a row of the frame d. The difference counts the
    respondents of both groups, which is the row of all respondents, or else
    the sum of the two groups."""
    if r["group"] != "difference":
        return r["n"], r["n_nonhan"]
    a = ame_get(d, group="all")
    if a is not None:
        return a["n"], a["n_nonhan"]
    parts = [ame_get(d, group=g) for g in ("agricultural", "non-agricultural")]
    if all(p is not None and num(p["n"]) is not None and num(p["n_nonhan"]) is not None for p in parts):
        return (str(sum(int(num(p["n"])) for p in parts)),
                str(sum(int(num(p["n_nonhan"])) for p in parts)))
    return r["n"], r["n_nonhan"]


#------------------------------------------------------------------------------
# Notes on merged birth cohorts, pooled provinces and merged categories
#------------------------------------------------------------------------------
def merges_of(note):
    """Merged birth cohorts in a note, as {first member: all members}"""
    out = {}
    m = re.search(r"birth cohorts merged: ([0-9+, ]+)", note or "")
    if m:
        for grp in m.group(1).split(","):
            ys = [int(y) for y in grp.strip().split("+") if y.strip()]
            if len(ys) > 1:
                out[ys[0]] = ys
    return out


def last_cohorts():
    """The latest birth cohort of each sample, from the coefficients and the
    notes. A merged group that holds it is the last group of its model."""
    mx = {}
    for f in ("S3_coef_A", "S3_coef_Aother", "S3_coef_B", "S3_coef_C"):
        d = rd(f)
        for s, t in zip(d["sample"], d["term"]):
            m = re.match(r"co_(\d+)", t)
            if m:
                mx[s] = max(mx.get(s, 0), int(m.group(1)))
    for f in AME_FILES:
        d = rd(f)
        for s, n in dict.fromkeys(zip(d["sample"], d["merge_note"])):
            for ys in merges_of(n).values():
                mx[s] = max(mx.get(s, 0), max(ys))
    return mx


LAST = {}


def cohort_span(ys, sample):
    """Birth years of a cohort group: 1935 stands for the years before 1940"""
    first, last = ys[0], max(ys)
    if LAST.get(sample) == last and len(ys) > 1:
        return f"{first} or later"
    if first == 1935:
        return f"before {last + 5}"
    return f"{first}{DASH}{last + 4}"


def merge_parts(note, sample, outcome):
    """The parts of a note in plain words"""
    out = []
    for part in (note or "").split("; "):
        if not part:
            continue
        if part.startswith("birth cohorts merged: "):
            for ys in merges_of(part).values():
                out.append(f"respondents born {cohort_span(ys, sample)} form one birth cohort group")
            continue
        m = re.match(r"provinces pooled into the small-province group: (.*)", part)
        if m:
            names = [PROVINCES.get(c.strip(), c.strip()) for c in m.group(1).split(",")]
            out.append(join_and(names) + " pooled into the small-province group")
            continue
        m = re.match(r"terms free across levels: (.*)", part)
        if m:
            out.append("coefficients free across the levels of education for "
                       + join_and(TERMS.get(t.strip(), t.strip()) for t in m.group(1).split(",")))
            continue
        if re.fullmatch(r"\d(\+\d)+", part):
            levels = EDU6_LEVELS if outcome == "edu6" else EDU_LEVELS
            out.append("the levels " + join_and(levels[int(k) - 1].lower() for k in part.split("+"))
                       + " merged into one category")
            continue
        if part == "two categories, logit" and out:
            out[-1] += ", which leaves two categories and a logit model"
            continue
        if part == "small-province group pooled with the reference province":
            out.append("the " + part)
            continue
        out.append(part)
    return out


def merge_text(note, sample, outcome=""):
    return "; ".join(merge_parts(note, sample, outcome))


#------------------------------------------------------------------------------
# Word tables, written as WordprocessingML so that long tables build fast
#------------------------------------------------------------------------------
def run_xml(text, size, bold=False, italic=False):
    rpr = ('<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial" w:eastAsia="Arial"/>'
           + ("<w:b/><w:bCs/>" if bold else "") + ("<w:i/><w:iCs/>" if italic else "")
           + f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>')
    return f'<w:r>{rpr}<w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def para_xml(text, size, bold=False, italic=False, indent=0, keep=False):
    ppr = ('<w:pPr>' + ("<w:keepNext/>" if keep else "")
           + '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
           + (f'<w:ind w:left="{indent}"/>' if indent else "") + '</w:pPr>')
    return f"<w:p>{ppr}{run_xml(text, size, bold, italic)}</w:p>"


def C(lines, span=1, bold=False, italic=False, indent=0):
    """A cell: one paragraph for each line"""
    if isinstance(lines, str):
        lines = [lines]
    return {"lines": lines, "span": span, "bold": bold, "italic": italic, "indent": indent}


def label_row(text, ncol, bold=True, italic=False):
    return [C(text, span=ncol, bold=bold, italic=italic)]


def table_xml(widths, head, body, size):
    """Three horizontal lines: above and below the heading rows, and below the
    last row. Heading rows repeat on every page."""
    rows = []
    allrows = [(r, True) for r in head] + [(r, False) for r in body]
    for i, (cells, is_head) in enumerate(allrows):
        top = i == 0
        bottom = (is_head and i == len(head) - 1) or i == len(allrows) - 1
        keep = i < len(allrows) - 1 and (is_head or (len(cells) == 1 and cells[0]["span"] > 1))
        trpr = "<w:trPr><w:cantSplit/>" + ("<w:tblHeader/>" if is_head else "") + "</w:trPr>"
        tcs, col = [], 0
        for c in cells:
            w = sum(widths[col:col + c["span"]])
            col += c["span"]
            b = (RULE.format("top") if top else "") + (RULE.format("bottom") if bottom else "")
            tcpr = (f'<w:tcPr><w:tcW w:w="{w}" w:type="dxa"/>'
                    + (f'<w:gridSpan w:val="{c["span"]}"/>' if c["span"] > 1 else "")
                    + (f"<w:tcBorders>{b}</w:tcBorders>" if b else "") + "</w:tcPr>")
            ps = "".join(para_xml(t, size, c["bold"] or is_head, c["italic"], c["indent"], keep)
                         for t in c["lines"])
            tcs.append(f"<w:tc>{tcpr}{ps}</w:tc>")
        assert col == len(widths), (col, len(widths), cells)
        rows.append(f"<w:tr>{trpr}{''.join(tcs)}</w:tr>")
    nil = "".join(f'<w:{s} w:val="nil"/>' for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
    tblpr = (f'<w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/>'
             f"<w:tblBorders>{nil}</w:tblBorders>"
             '<w:tblLayout w:type="fixed"/>'
             '<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/>'
             '<w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar>'
             '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" '
             'w:lastColumn="0" w:noHBand="1" w:noVBand="1"/></w:tblPr>')
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
    return f"<w:tbl {nsdecls('w')}>{tblpr}<w:tblGrid>{grid}</w:tblGrid>{''.join(rows)}</w:tbl>"


def add_table(doc, widths, head, body, size=22):
    """Main tables in 11 point, supplementary tables in 8 point (size=16)"""
    assert abs(sum(widths) - TEXT_WIDTH) <= 5, sum(widths)
    body_el = doc.element.body
    body_el.insert(len(body_el) - 1, parse_xml(table_xml(widths, head, body, size)))
    return len(head) + len(body)


#------------------------------------------------------------------------------
# Paragraphs
#------------------------------------------------------------------------------
def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    pf = st.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    return doc


def para(doc, parts, size=11, after=0, before=0, keep=False, page_break=False):
    """parts is a list of (text, bold, italic)"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
    pf.keep_with_next = keep
    pf.page_break_before = page_break
    for text, b, i in parts:
        r = p.add_run(text)
        r.bold, r.italic = b, i
        r.font.name, r.font.size = "Arial", Pt(size)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    return p


def caption(doc, label, title, page_break=True):
    para(doc, [(label, True, False)], keep=True, page_break=page_break)
    para(doc, [(title, False, True)], keep=True, after=6)


def notes(doc, items, size=11):
    """Each note is one paragraph that starts with a bold Note:. A word set
    between asterisks is printed in italics."""
    for k, text in enumerate(items):
        parts = [("Note: ", True, False)]
        for bit in re.split(r"(\*[^*]+\*)", text):
            if bit:
                parts.append((bit.strip("*"), False, bit.startswith("*")))
        para(doc, parts, size=size, before=6 if k == 0 else 3)


P_NOTE = "The ( ) includes the *p*-value and the [ ] the 95% confidence interval."
COUNT_NOTE = ("<10 marks a count from 1 to 9, and suppressed marks a count hidden because it could be "
              "used to calculate a count below 10.")
SUPPRESS_NOTE = (f"<10 marks a count from 1 to 9, suppressed marks a count hidden because it could be "
                 f"used to calculate a count below 10, and {DASH} marks the estimate of such a cell.")


def count_note(*frames):
    """The note on hidden counts, for a listing that holds one"""
    for d in frames:
        for c in ("n", "n_nonhan"):
            if c in d.columns and d[c].isin(HIDDEN).any():
                return [COUNT_NOTE]
    return []


#------------------------------------------------------------------------------
# Main tables
#------------------------------------------------------------------------------
def table_1(doc):
    t1 = rd("T1_describe")
    cols = ["Han", "Non-Han", "Han, agricultural", "Non-Han, agricultural",
            "Han, non-agricultural", "Non-Han, non-agricultural"]
    rows_def = [
        ("Hukou at age 3", "Agricultural", "Agricultural hukou at age 3 (%)", 1),
        ("Current hukou", "Agricultural", "Agricultural hukou at the interview (%)", 1),
        ("Gender", "Male", "Male (%)", 1),
        ("Age", "Mean", "Age, mean (years)", 1),
        ("Parents' education", "Mean level, 1 to 5", "Parents' education, mean level (1 to 5)", 2),
        ("Education", None, "Education (%)", 1),
        ("Years of schooling", "Mean", "Years of schooling, mean", 1),
        ("Home language", None, "Home language (%)", 1),
        ("Mandarin rating", "Mean, 1 to 7", "Interviewer's rating of Mandarin, mean (1 to 7)", 2),
        ("Word test", "Mean", "Word test score, mean (points)", 1),
        ("Respondents", "n", "Respondents (n)", None)]
    widths = [2520] + [1140] * 6
    head = [[C(""), C("All", 2), C("Agricultural hukou at age 3", 2), C("Non-agricultural hukou at age 3", 2)],
            [C("")] + [C(x) for x in ["Han", "Non-Han"] * 3]]
    body = []
    for s, year in (("A-2010", "2010"), ("A-2018", "2018")):
        d = t1[t1["sample"] == s]
        body.append(label_row(f"CFPS {year}", 7))
        for var, lev, lab, dec in rows_def:
            dv = d[d["variable"] == var]
            if dv.empty:
                continue
            levels = [lev] if lev else list(dict.fromkeys(dv["level"]))
            if not lev:
                body.append(label_row(lab, 7, bold=False))
            for lv in levels:
                cells = [C(lab if lev else lv, indent=0 if lev else INDENT)]
                for g in cols:
                    r = ame_get(dv, level=lv, group=g)
                    if r is None:
                        cells.append(C(DASH))
                    elif dec is None:
                        cells.append(C(fmt_n(r["n_shown"])))
                    else:
                        cells.append(C(fmt(r["estimate"], dec)))
                body.append(cells)
    caption(doc, "Table I", "Characteristics of the Samples by Ethnicity and Hukou at Age 3, CFPS 2010 and 2018",
            page_break=False)
    n = add_table(doc, widths, head, body)
    notes(doc, [
        "Using CFPS 2010 and 2018 data. Percentages and means are weighted with the CFPS national "
        "cross-sectional weight of each wave and the survey design, and the number of respondents is "
        "unweighted. Each sample holds the respondents aged 25 or older with a positive weight whose "
        "ethnicity, hukou at age 3, birth year, gender, province and education are known (Table S1).",
        f"{DASH} marks a value that is not reported because the cell holds fewer than 10 respondents, or "
        "because the count of the cell could be calculated from the other cells. Home language and the "
        "interviewer's rating of Mandarin are shown for 2010 only."])
    return n


def table_2(doc):
    t2 = rd("T2_ame_main")
    models = ("M1", "M2", "M3", "M4")
    widths = [2360] + [1750] * 4
    head = [[C("")] + [C(m) for m in models]]
    body = []
    for s, year in (("A-2010", "2010"), ("A-2018", "2018")):
        d = t2[t2["sample"] == s]
        body.append(label_row(f"CFPS {year}", 5))
        for tg in ("ge3", "ge4", "ge5"):
            body.append(label_row(TARGETS["edu5"][tg], 5, bold=False, italic=True))
            for g, glab in GROUPS3:
                body.append([C(glab, indent=INDENT)]
                            + [C(ame_lines(ame_get(d, model=m, target=tg, group=g))) for m in models])
        body.append(label_row("Respondents, n (non-Han n)", 5, bold=False, italic=True))
        for g, glab in GROUPS4[:3]:
            body.append([C(glab, indent=INDENT)]
                        + [C(n_text(ame_get(d, model=m, target="ge3", group=g))) for m in models])
        body.append([C("Imputed data sets")]
                    + [C("None" if ame_get(d, model=m, target="ge3", group="all")["m"] == "1"
                         else ame_get(d, model=m, target="ge3", group="all")["m"]) for m in models])
    caption(doc, "Table II", "Average Marginal Effects of Non-Han Ethnicity on Educational Attainment by Hukou "
                             "at Age 3, CFPS 2010 and 2018")
    n = add_table(doc, widths, head, body)
    r18 = ame_get(t2, sample="A-2018", model="M4", target="ge3", group="all")
    merge18 = merge_text(r18["merge_note"], "A-2018", "edu5")
    notes(doc, [
        "Each cell shows the average marginal effect of non-Han ethnicity, in percentage points, on the "
        "probability of reaching the stated level of education, within the group of hukou at age 3. The "
        "difference is the effect in the agricultural hukou group minus the effect in the non-agricultural "
        "hukou group. " + P_NOTE,
        "Ordered logit models of five education levels, weighted with the CFPS national cross-sectional "
        "weights and the survey design. M1 holds non-Han ethnicity, birth cohort and gender, so its two "
        "group effects differ only through the composition of the groups. M2 adds hukou at age 3 and its "
        "interaction with non-Han ethnicity, M3 adds parents' education and M4 adds province. M3 and M4 "
        "combine 20 imputed data sets of parents' education by Rubin's rules.",
        "A birth cohort group with fewer than 30 respondents, or whose respondents are all in the lowest or "
        "all in the highest education level, is merged with a neighbouring group"
        + (f", and in the CFPS 2018 models {merge18}" if merge18 else "")
        + ". Table S3 reports every coefficient and the effects on primary school or above and on the "
        "expected education level."])
    return n


def table_3(doc):
    t3 = rd("T3_ame_language")
    widths = [3160, 2200, 2200, 1800]
    head = [[C(""), C("M4"), C("M4 + education"), C("n (non-Han n)")]]
    body = []
    blocks = [
        ("Home language (probability, percentage points)", "lang",
         [("C-lang-2010", "2010"), ("C-lang-2012", "2012")], ["pr_mandarin", "pr_minority", "pr_dialect"]),
        ("Interviewer's rating of Mandarin (expected rating, 1 to 7)", "rating",
         [("C-rating-2010", "2010"), ("C-rating-2012", "2012"), ("C-rating-2014", "2014"),
          ("C-rating-2016", "2016")], ["mean"]),
        ("Word test score (expected score, points)", "wordtest",
         [("C-wordtest-2010", "2010"), ("C-wordtest-2014", "2014"), ("C-wordtest-2018", "2018")], ["mean"])]
    lang_short = {"pr_mandarin": "Mandarin", "pr_minority": "minority language", "pr_dialect": "Chinese dialect"}
    oc_short = {"lang": "home language", "rating": "rating", "wordtest": "word test"}
    extra, merged = [], {}
    for title, oc, samples, tgs in blocks:
        body.append(label_row(title, 4))
        for s, year in samples:
            d = t3[t3["sample"] == s]
            if (d["target"] == "not estimable").all():
                r = d.iloc[0]
                body.append([C(f"{year}, not estimable", italic=True), C(DASH), C(DASH), C(n_text(r))])
                extra.append(f"The {year} {oc_short[oc]} model is not estimable because {r['merge_note']}.")
                continue
            parts = []
            for nt in dict.fromkeys(d["merge_note"]):
                parts += [p for p in merge_parts(nt, s, oc) if p not in parts]
            if parts:
                merged.setdefault((year, join_and(parts)), []).append(oc_short[oc])
            for tg in tgs:
                body.append(label_row(year if len(tgs) == 1 else f"{year}, {lang_short[tg]}", 4,
                                      bold=False, italic=True))
                for g, glab in GROUPS3:
                    r4 = ame_get(d, model="M4", target=tg, group=g)
                    r5 = ame_get(d, model="M4 + education", target=tg, group=g)
                    rn = ame_get(d, model="M4", target=tg, group="all" if g == "difference" else g)
                    body.append([C(glab, indent=INDENT), C(ame_lines(r4)), C(ame_lines(r5)), C(n_text(rn))])
    # Models of one year with the same merges share one sentence
    extra += [f"In the {year} {join_and(ocs)} models, {txt}." for (year, txt), ocs in merged.items()]
    caption(doc, "Table III", "Average Marginal Effects of Non-Han Ethnicity on Home Language, the Interviewer's "
                              "Rating of Mandarin and the Word Test Score by Hukou at Age 3")
    n = add_table(doc, widths, head, body)
    notes(doc, [
        "Each cell shows the average marginal effect of non-Han ethnicity within the group of hukou at age 3, "
        "in percentage points for home language and in points of the outcome for the rating and the word "
        "test. The difference is the effect in the agricultural hukou group minus the effect in the "
        "non-agricultural hukou group. n (non-Han n) counts the respondents of the group, and of both groups "
        "for the difference. " + P_NOTE,
        "Multinomial logit models of home language with Chinese dialect as the base, ordered logit models "
        "of the rating and linear regressions of the word test score, weighted with the CFPS national "
        "cross-sectional weights and the survey design. Each sample holds the respondents aged 16 or older "
        "with a positive weight whose outcome, ethnicity, hukou at age 3, birth year, gender and province "
        "are known (Table S1). M4 holds the terms of M4 in Table II, and M4 + education adds the "
        "respondent's education. Both combine 20 imputed data sets by Rubin's rules.",
        " ".join(extra) + " Table S3 reports every coefficient and the effects on each rating category."])
    return n


def table_4(doc):
    k = rd("T4_khb")
    b = rd("T4_bcohort")
    widths = [3360, 2000, 2000, 2000]
    head = [[C("")] + [C(glab) for _, glab in GROUPS3]]
    body = [label_row("A. Decomposition of the non-Han coefficient, five education levels, CFPS 2010", 4)]
    steps = [("step 1", "Step 1, home language"),
             ("step 2", "Step 2, home language, rating of Mandarin and word test")]
    for st, stlab in steps:
        body.append(label_row(stlab, 4, bold=False, italic=True))
        d = k[k["step"] == st]
        for q, qlab in KHB.items():
            if q == "contrib_language" or (d["quantity"] == q).sum() == 0:
                continue
            cells = [C(qlab, indent=2 * INDENT if q.startswith("contrib") else INDENT)]
            for g, _ in GROUPS3:
                r = ame_get(d, quantity=q, group=g)
                if r is None or num(r["estimate"]) is None:
                    cells.append(C(DASH))
                elif q == "pct_indirect":
                    cells.append(C(fmt(r["estimate"], 1)))
                else:
                    cells.append(C([f"{fmt(r['estimate'], 2)} ({fmt_p(r['p'])})",
                                    f"[{fmt(r['ci_low'], 2)}, {fmt(r['ci_high'], 2)}]"]))
            body.append(cells)
    d = k[(k["step"] == "step 1") & (k["quantity"] == "total")]
    body.append([C("n (non-Han n)")]
                + [C("{} ({})".format(*map(fmt_n, counts_of(d, ame_get(d, group=g))))) for g, _ in GROUPS3])
    body.append(label_row("B. Children aged 12 to 15 in 2010, education by 2018 (percentage points)", 4))
    models = [("B-cohort", "base", "Base model"),
              ("B-cohort-lang", "base, language sample", "Base model, language sample"),
              ("B-cohort-lang", "base + language", "Base model + language measures")]
    for oc in ("y_sh", "y_ter"):
        body.append(label_row(TARGETS[oc]["pr"], 4, bold=False, italic=True))
        for s, m, mlab in models:
            body.append([C(mlab, indent=INDENT)]
                        + [C(ame_lines(ame_get(b, sample=s, model=m, outcome=oc, group=g))) for g, _ in GROUPS3])
    body.append(label_row("n (non-Han n), all children", 4, bold=False, italic=True))
    for s, m, lab in (("B-cohort", "base", "Base model"),
                      ("B-cohort-lang", "base, language sample", "Language sample")):
        body.append([C(lab, indent=INDENT),
                     C(n_text(ame_get(b, sample=s, model=m, outcome="y_ter", group="all")), span=3)])
    caption(doc, "Table IV", "Decomposition of the Non-Han Gap in Education and Education by 2018 of the "
                             "Children Aged 12 to 15 in 2010")
    n = add_table(doc, widths, head, body)
    rs = ame_get(b, sample="B-cohort", model="base", outcome="y_sh", group="agricultural")
    pooled = [p for p in merge_parts(rs["merge_note"], "B-cohort", "y_sh") if "pooled" in p]
    notes(doc, [
        "Panel A decomposes the coefficient of non-Han ethnicity in the ordered logit model M4 of five "
        "education levels (CFPS 2010) into a direct part and a part through the language measures, on the "
        "scale of the ordered logit coefficients, by the method of Karlson, Holm and Breen. Step 1 uses home "
        "language, and step 2 adds the interviewer's rating of Mandarin and the word test score. The rating "
        "and the word test were measured at the interview, after most respondents had left school, so "
        "schooling can affect both. The sample holds the respondents of the CFPS 2010 sample of Table II "
        "whose three language measures are known, and n (non-Han n) counts the respondents of the group, "
        "and of both groups for the difference. Standard errors include the estimation of every model in "
        "the decomposition and combine 20 imputed data sets by Rubin's rules.",
        "Panel B shows average marginal effects of non-Han ethnicity, in percentage points, from logit "
        "models of the children aged 12 to 15 in 2010 who were interviewed again in 2018, weighted with the "
        "CFPS 2010 to 2018 panel weight and the survey design. The base model holds non-Han ethnicity, "
        "hukou at age 3 and their interaction, gender, parents' education, log family income per person, "
        "province and age in 2010, and the language measures are home language, the interviewer's rating "
        "of Mandarin and the word test score of 2010. The language sample holds the children whose three "
        "language measures are known. Each model combines 20 imputed data sets by Rubin's rules.",
        f"{DASH} marks an effect that is not reported because fewer than 10 non-Han children had "
        "non-agricultural hukou at age 3. For senior high school or above, the outcome does not vary among "
        "the non-Han children with non-agricultural hukou at age 3, so these models were fitted without them "
        "and without the interaction, and only the effect for the agricultural hukou group is shown"
        + (f", and {pooled[0].replace(' pooled', ' was pooled')}" if pooled else "") + ".",
        "In both panels, the difference is the effect in the agricultural hukou group minus the effect in "
        "the non-agricultural hukou group. " + P_NOTE])
    return n


#------------------------------------------------------------------------------
# Figure 1
#------------------------------------------------------------------------------
def pick_font():
    from matplotlib import font_manager
    names = {f.name for f in font_manager.fontManager.ttflist}
    return "Arial" if "Arial" in names else "DejaVu Sans"


def figure_1():
    f1 = rd("F1_transitions")
    cohorts = list(dict.fromkeys(f1["cohort"]))
    xlab = [c.replace("-", DASH) for c in cohorts]
    src_of = {c: s for c, s in zip(f1["cohort"], f1["source"])}
    split = max(i for i, c in enumerate(cohorts) if src_of[c] == "A-2010") + 0.5
    series = [
        ("Han, agricultural", "Han, agricultural hukou at age 3", "0.55", "-", "o", True),
        ("Han, non-agricultural", "Han, non-agricultural hukou at age 3", "0.55", "--", "o", False),
        ("Non-Han, agricultural", "Non-Han, agricultural hukou at age 3", "0", "-", "s", True),
        ("Non-Han, non-agricultural", "Non-Han, non-agricultural hukou at age 3", "0", "--", "s", False)]
    panels = [("T1 Primary school", "T1, primary school among all"),
              ("T2 Junior high school", "T2, junior high school among those\nwith primary school"),
              ("T3 Senior high school", "T3, senior high school among those\nwith junior high school"),
              ("T4 College", "T4, college among those with\nsenior high school")]
    plt.rcParams.update({"font.family": pick_font(), "font.size": 8, "axes.linewidth": 0.6,
                         "xtick.major.width": 0.6, "ytick.major.width": 0.6, "pdf.fonttype": 42})
    fig, axes = plt.subplots(2, 2, figsize=(6.5, 6.9), sharex=True, sharey=True)
    shown = 0
    for ax, (tr, title) in zip(axes.flat, panels):
        for key, _, grey, ls, mk, filled in series:
            for src in ("A-2010", "A-2018"):
                xs, ys = [], []
                for i, c in enumerate(cohorts):
                    if src_of[c] != src:
                        continue
                    r = f1[(f1["cohort"] == c) & (f1["group"] == key) & (f1["transition"] == tr)]
                    v = num(r.iloc[0]["estimate"]) if not r.empty else None
                    xs.append(i)
                    ys.append(float("nan") if v is None else v)
                    shown += v is not None
                ax.plot(xs, ys, color=grey, linestyle=ls, linewidth=1.1, marker=mk, markersize=3.8,
                        markerfacecolor=grey if filled else "white", markeredgecolor=grey,
                        markeredgewidth=0.9, solid_capstyle="round", dash_capstyle="round", zorder=3)
        ax.axvline(split, color="0", linewidth=0.5, zorder=1)
        ax.set_title(title, fontsize=8, loc="left", pad=4)
        ax.set_ylim(0, 100)
        ax.set_yticks([0, 25, 50, 75, 100])
        ax.grid(axis="y", color="0.85", linewidth=0.4, zorder=0)
        ax.set_axisbelow(True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        ax.set_xticks(range(len(cohorts)))
        ax.set_xticklabels(xlab, rotation=60, ha="right", rotation_mode="anchor", fontsize=7)
        ax.tick_params(axis="y", labelsize=7)
    a0 = axes[0, 0]
    a0.text(split - 0.25, 3, "CFPS 2010", ha="right", va="bottom", fontsize=6.5)
    a0.text(split + 0.25, 3, "CFPS 2018", ha="left", va="bottom", fontsize=6.5)
    for ax in axes[:, 0]:
        ax.set_ylabel("Percent", fontsize=8)
    for ax in axes[1, :]:
        ax.set_xlabel("Birth cohort", fontsize=8)
    handles = [Line2D([0], [0], color=g, linestyle=ls, linewidth=1.1, marker=mk, markersize=4.2,
                      markerfacecolor=g if fl else "white", markeredgecolor=g, markeredgewidth=0.9)
               for _, _, g, ls, mk, fl in series]
    fig.legend(handles, [lab for _, lab, *_ in series], loc="lower center", ncol=2, frameon=False,
               fontsize=7.5, handlelength=3.2, columnspacing=2.0, bbox_to_anchor=(0.5, 0.0))
    fig.tight_layout(rect=(0, 0.07, 1, 1), h_pad=1.6, w_pad=1.2)
    fig.savefig(os.path.join(OUT, "Figure 1.png"), dpi=600, facecolor="white")
    fig.savefig(os.path.join(OUT, "Figure 1.pdf"), facecolor="white")
    plt.close(fig)
    return shown


def figure_page(doc):
    f1 = rd("F1_transitions")
    gone = []
    for g, d in f1.groupby("group", sort=False):
        k = int((d["estimate"] == "").sum())
        if k:
            eth, huk = g.split(", ")
            gone.append(f"{k} of the {len(d)} points of the {eth.replace('Non-Han', 'non-Han')} group with "
                        f"{huk} hukou at age 3")
    caption(doc, "Figure 1", "Share Completing Each School Transition by Birth Cohort, Ethnicity and Hukou "
                             "at Age 3")
    doc.add_picture(os.path.join(OUT, "Figure 1.png"), width=Inches(6.5))
    notes(doc, [
        "Using CFPS 2010 and 2018 data. Percentages are weighted with the CFPS national cross-sectional "
        "weights of each wave. Cohorts born up to 1985 come from CFPS 2010 and cohorts born from 1986 to "
        "1993 from CFPS 2018, so each line breaks where the source changes, at the vertical line.",
        "T1 is completing primary school among all respondents, T2 completing junior high school among "
        "those who completed primary school, T3 completing senior high school among those who completed "
        "junior high school and T4 completing college or above among those who completed senior high "
        "school. A point based on fewer than 10 respondents is not shown"
        + (", which removes " + join_and(gone) if gone else "") + "."])


#------------------------------------------------------------------------------
# Supplementary tables
#------------------------------------------------------------------------------
AME_W = [2660, 1000, 900, 1700, 900, 1100, 1100]
AME_HEAD = [[C(x) for x in ("", "Estimate", "S.E.", "95% C.I.", "p-value", "n", "Non-Han n")]]
COEF_W = [4160, 1300, 1200, 1800, 900]
COEF_HEAD = [[C(x) for x in ("", "Estimate", "S.E.", "95% C.I.", "p-value")]]


def model_heading(r, prefix=""):
    imp = "" if r["m"] in ("", "1") else f", {r['m']} imputed data sets"
    mt = merge_text(r.get("merge_note", ""), r["sample"], r["outcome"])
    return (f"{prefix}{r['sample']}, {r['model']}, {OUTCOMES.get(r['outcome'], r['outcome'])}{imp}"
            + (f" ({mt})" if mt else ""))


def ame_rows(d, prefix=""):
    """Rows of an AME listing: a heading for each model, a subheading for
    each target and one row for each group"""
    body = []
    for s, m, oc in dict.fromkeys(zip(d["sample"], d["model"], d["outcome"])):
        dm = d[(d["sample"] == s) & (d["model"] == m) & (d["outcome"] == oc)]
        r0 = dm.iloc[0]
        body.append(label_row(model_heading(r0, prefix), 7))
        if (dm["target"] == "not estimable").all():
            body.append([C("Not estimable", italic=True, indent=INDENT)] + [C(DASH)] * 4
                        + [C(fmt_n(r0["n"])), C(fmt_n(r0["n_nonhan"]))])
            continue
        for tg in dict.fromkeys(dm["target"]):
            dt = dm[dm["target"] == tg]
            if tg == "pct_indirect":
                tlab = KHB[tg]
            elif tg in KHB:
                tlab = KHB[tg] + ", ordered logit coefficient"
            else:
                tlab = TARGETS.get(oc, {}).get(tg, tg) + (", percentage points" if is_prob(oc, tg) else "")
            body.append(label_row(tlab, 7, bold=False, italic=True))
            for g, glab in GROUPS4:
                r = ame_get(dt, group=g)
                if r is None:
                    continue
                if is_prob(oc, tg):
                    sc, dd = 100.0, 2
                elif tg == "pct_indirect":
                    sc, dd = 1.0, 1
                else:
                    sc, dd = 1.0, 3
                ratio = tg == "pct_indirect"
                est = num(r["ci_low"]) is not None and not ratio
                n, nn = counts_of(dt, r)
                body.append([C(glab, indent=2 * INDENT),
                             C(fmt(r["estimate"], dd, sc)), C(DASH if ratio else fmt(r["se"], dd, sc)),
                             C(f"[{fmt(r['ci_low'], dd, sc)}, {fmt(r['ci_high'], dd, sc)}]" if est else DASH),
                             C(DASH if ratio else fmt_p(r["p"])), C(fmt_n(n)), C(fmt_n(nn))])
    return body


AME_NOTES = [
    "Average marginal effects of non-Han ethnicity, in percentage points for a probability and in points "
    "of the outcome for an expected value, for all respondents and within each group of hukou at age 3. "
    "The difference is the effect in the agricultural hukou group minus the effect in the non-agricultural "
    "hukou group, and its n and non-Han n count the respondents of both groups. Standard errors use the "
    "survey design and combine the imputed data sets by Rubin's rules.",
    f"{DASH} marks an estimate that is not reported, because the group holds fewer than 10 non-Han "
    "respondents or because the model was fitted without that group. A birth cohort group with fewer than "
    "30 respondents, or whose respondents are all in the lowest or all in the highest outcome category, is "
    "merged with a neighbouring group, and an outcome category with fewer than 20 Han or 20 non-Han "
    "respondents is merged with a neighbouring category. The heading of each model states such merges."]

COEF_NOTES = [
    "Coefficients on the scale of the model (log odds for the logit and ordered logit models, points of the "
    "outcome for the linear regressions), with standard errors from the survey design combined across the "
    "imputed data sets by Rubin's rules. Cut points of the ordered logit models are on the same scale.",
    "The reference categories are Han ethnicity, non-agricultural hukou at age 3, female, the earliest "
    "birth cohort group of each model, parents with illiterate or semi-literate education, the largest "
    "province group of each model and, where present, illiterate or semi-literate education, the youngest "
    "age and the wave 2010. A province with fewer than 30 respondents, or whose respondents are all in the "
    "lowest or all in the highest outcome category, joins the small-province group."]


def term_label(term, sample, merges):
    eq = ""
    if term.startswith("(Intercept):"):
        term, eq = "(Intercept)", term.split(":", 1)[1]
    elif ":" in term:
        term, eq = term.split(":", 1)
    fixed = {
        "(Intercept)": "Intercept", "nonhan": "Non-Han", "agri3": "Agricultural hukou at age 3",
        "nh_ag": "Non-Han × agricultural hukou at age 3",
        "nonhan + nh_ag": "Non-Han in the agricultural hukou group, non-Han plus the interaction",
        "male": "Male", "linc": "Log family income per person", "rating": "Interviewer's rating of Mandarin",
        "wordtest": "Word test score", "lang_min": "Home language, minority language",
        "lang_dia": "Home language, Chinese dialect", "pv_0": "Province, small-province group"}
    if term in fixed:
        lab = fixed[term]
    elif re.fullmatch(r"\d\|\d", term):
        lab = f"Cut point {term}"
    elif term.startswith("pe_"):
        lab = "Parents' education, " + EDU_LEVELS[int(term[3:]) - 1].lower()
    elif term.startswith("ed_"):
        lab = "Education, " + EDU_LEVELS[int(term[3:]) - 1].lower()
    elif term.startswith("pv_"):
        lab = "Province, " + PROVINCES.get(term[3:], term[3:])
    elif term.startswith("co_"):
        y = int(term[3:])
        lab = "Born " + cohort_span(merges.get(y, [y]), sample)
    elif term.startswith("ag_"):
        lab = f"Age {term[3:]}"
    elif term.startswith("wv_"):
        lab = f"Wave 20{term[3:]}"
    else:
        lab = term
    return lab, eq


def coef_rows(d, ame_tables):
    """Rows of a coefficient listing with a heading for each model. The
    notes come from the AME tables of the same models."""
    notes_of = {}
    for t in ame_tables:
        for s, m, oc, nt in zip(t["sample"], t["model"], t["outcome"], t["merge_note"]):
            notes_of.setdefault((s, m, oc), nt)
    eq_names = {"Mandarin": "Mandarin", "Minority": "Minority language"}
    body = []
    for s, m, oc in dict.fromkeys(zip(d["sample"], d["model"], d["outcome"])):
        dm = d[(d["sample"] == s) & (d["model"] == m) & (d["outcome"] == oc)]
        r0 = dm.iloc[0]
        note = notes_of.get((s, m, oc), "")
        imp = "" if r0["m"] in ("", "1") else f", {r0['m']} imputed data sets"
        mt = merge_text(note, s, oc)
        body.append(label_row(f"{s}, {m}, {OUTCOMES.get(oc, oc)}, n = {fmt_n(r0['n'])}, non-Han n = "
                              f"{fmt_n(r0['n_nonhan'])}{imp}" + (f" ({mt})" if mt else ""), 5))
        merges = merges_of(note)
        rows = [(term_label(t, s, merges), r) for t, r in zip(dm["term"], dm.to_dict("records"))]
        for eq in dict.fromkeys(e for (_, e), _ in rows):
            if eq:
                body.append(label_row(f"{eq_names.get(eq, eq)} versus Chinese dialect", 5, bold=False,
                                      italic=True))
            for (lab, e), r in rows:
                if e != eq:
                    continue
                est = num(r["ci_low"]) is not None
                body.append([C(lab, indent=2 * INDENT if eq else INDENT), C(fmt(r["estimate"], 3)),
                             C(fmt(r["se"], 3)),
                             C(f"[{fmt(r['ci_low'], 3)}, {fmt(r['ci_high'], 3)}]" if est else DASH),
                             C(fmt_p(r["p"]))])
    return body


def continued(doc, label, title):
    para(doc, [(f"{label} (continued)", True, False)], keep=True, page_break=True)
    para(doc, [(title, False, True)], keep=True, after=6)


def supp_flow(doc):
    s1 = rd("S1_flow")
    groups = ("Han", "Non-Han", "Ethnicity missing", "All")
    widths = [4560, 1200, 1200, 1200, 1200]
    head = [[C(x) for x in ("",) + groups]]
    body = []
    for s in dict.fromkeys(s1["sample"]):
        d = s1[s1["sample"] == s]
        body.append(label_row(s, 5))
        for st in dict.fromkeys(d["step"]):
            ds = d[d["step"] == st]
            body.append([C(st, indent=INDENT)]
                        + [C(DASH if ame_get(ds, group=g) is None else fmt_n(ame_get(ds, group=g)["n_shown"]))
                           for g in groups])
    caption(doc, "Table S1", "Sample Flow, Respondents Kept at Each Step of Each Sample", page_break=False)
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Unweighted numbers of respondents. Each step keeps the respondents of the step above who meet the "
        "stated condition, so the number left out at a step is the difference between the two steps.",
        "Samples A hold the adults aged 25 or older (Tables I, II, S4 and S5), samples C the adults aged 16 "
        "or older (Table III) and sample D-A10 the respondents of A-2010 whose three language measures are "
        "known (Table IV, panel A). Samples B-cohort and B-cohort-lang hold the children aged 12 to 15 in "
        "2010 who were interviewed again in 2018, all of them and those whose three language measures of "
        "2010 are known (Table IV, panel B), and samples B-cohort-10to15 and B-cohort-lang-10to15 the same "
        "for the children aged 10 to 15 in 2010 (check R14). Sample B-enrol-16-22 holds the respondents aged "
        "16 to 22 of the waves 2010 to 2018, pooled, and samples B-enrol-16-18 and B-enrol-19-22 its two age "
        "groups (Table S6). Samples R15 hold the samples of check R15. The last panel, Codes set to missing, "
        "counts the respondents whose answer was in a code that the analysis sets to missing.",
        COUNT_NOTE], size=9)
    return n


def supp_groups(doc):
    g = rd("S2_groups")
    p = rd("S2_provinces")
    samples = list(dict.fromkeys(g["sample"]))
    widths = [1560] + [1080] * (len(samples) - 1) + [1320]
    head = [[C("")] + [C(s) for s in samples]]
    body = [label_row("Non-Han respondents by ethnic group, n (% of the non-Han respondents)", 8)]
    for name in dict.fromkeys(g["name"]):
        cells = [C(name, indent=INDENT)]
        for s in samples:
            r = ame_get(g, sample=s, name=name)
            if r is None:
                cells.append(C(DASH))
            else:
                pct = fmt(r["pct"], 1)
                cells.append(C(fmt_n(r["n_shown"]) + ("" if pct == DASH else f" ({pct})")))
        body.append(cells)
    for grp in ("Han", "Non-Han"):
        body.append(label_row(f"{grp} respondents by province, n", 8))
        for code in sorted(set(p.loc[p["group"] == grp, "prov_code"]), key=int):
            cells = [C(PROVINCES.get(code, code), indent=INDENT)]
            for s in samples:
                r = ame_get(p, sample=s, group=grp, prov_code=code)
                cells.append(C(DASH if r is None else fmt_n(r["n_shown"])))
            body.append(cells)
    caption(doc, "Table S2", "Non-Han Ethnic Groups and Provinces in the Samples")
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Unweighted numbers of respondents in the analysis samples. Ethnic groups are those recorded by CFPS "
        f"with at least one respondent in a sample, and {DASH} marks a group or province with no respondent "
        "in that sample.", COUNT_NOTE], size=9)
    return n


def supp_full(doc):
    t2, t3, t4b = rd("T2_ame_main"), rd("T3_ame_language"), rd("T4_bcohort")
    k = rd("T4_khb")
    body = [label_row("A. Average marginal effects of Tables II, III and IV", 7)]
    body += ame_rows(pd.concat([t2, t3, t4b], ignore_index=True))
    body.append(label_row("Table IV, panel A, decomposition of the non-Han coefficient, D-A10, five education "
                          "levels, ordered logit, 20 imputed data sets", 7))
    for st in dict.fromkeys(k["step"]):
        ks = k[k["step"] == st]
        for g, glab in GROUPS3:
            body.append(label_row(f"{cap(st)}, {glab[0].lower() + glab[1:]}", 7, bold=False, italic=True))
            for _, r in ks[ks["group"] == g].iterrows():
                ratio = r["quantity"] == "pct_indirect"
                dd = 1 if ratio else 3
                est = num(r["ci_low"]) is not None and not ratio
                n, nn = counts_of(ks[ks["quantity"] == r["quantity"]], r)
                body.append([C(KHB.get(r["quantity"], r["quantity"]), indent=2 * INDENT),
                             C(fmt(r["estimate"], dd)), C(DASH if ratio else fmt(r["se"], dd)),
                             C(f"[{fmt(r['ci_low'], dd)}, {fmt(r['ci_high'], dd)}]" if est else DASH),
                             C(DASH if ratio else fmt_p(r["p"])), C(fmt_n(n)), C(fmt_n(nn))])
    caption(doc, "Table S3", "Full Results of the Models in Tables II to IV")
    n = add_table(doc, AME_W, AME_HEAD, body, size=16)
    notes(doc, AME_NOTES + [
        "In the decomposition, estimates are on the scale of the ordered logit coefficients, and the "
        "contribution of home language through both categories is the sum of the contributions of minority "
        "language and Chinese dialect. The indirect part as a percentage of the total is shown without a "
        "standard error, as in Table IV, because a ratio has no reliable standard error when the total is "
        "close to zero."] + count_note(t2, t3, t4b, k), size=9)
    a, c, bco = rd("S3_coef_A"), rd("S3_coef_C"), rd("S3_coef_B")
    body = [label_row("B. Coefficients of the models in Tables II, III and IV", 5)]
    body += coef_rows(a[a["outcome"] == "edu5"], (t2,))
    body += coef_rows(c, (t3,))
    body += coef_rows(bco[bco["sample"].str.startswith("B-cohort")], (t4b,))
    continued(doc, "Table S3", "Coefficients")
    n += add_table(doc, COEF_W, COEF_HEAD, body, size=16)
    notes(doc, COEF_NOTES, size=9)
    return n


def supp_ame(doc, label, title, frame, extra_notes, coef):
    caption(doc, label, title)
    n = add_table(doc, AME_W, AME_HEAD, ame_rows(frame), size=16)
    notes(doc, AME_NOTES + list(extra_notes) + count_note(frame), size=9)
    continued(doc, label, "Coefficients")
    n += add_table(doc, COEF_W, COEF_HEAD, coef_rows(coef, (frame,)), size=16)
    notes(doc, COEF_NOTES + count_note(coef), size=9)
    return n


CHECKS = {
    "R1": "R1, no weights", "R2": "R2, complete cases without imputation",
    "R3": "R3, proxy interviews excluded",
    "R5": "R5, each of the largest non-Han groups compared with the Han respondents",
    "R6": "R6, models by birth cohort",
    "R7": "R7, hukou at age 3 only, without the values taken from hukou at age 12",
    "R9": "R9, province of 2010 in place of the province of the wave",
    "R10": "R10, partial proportional odds models",
    "R11": "R11, national resample with its own weights",
    "R12": "R12, stages by the highest stage attended",
    "R13": "R13, six education levels",
    "R14": "R14, children aged 10 to 15 in 2010",
    "R15": "R15, code 79 of the rating of Mandarin as the lowest category"}


def supp_checks(doc):
    r = rd("S7_robust")
    body = []
    for ck in sorted(dict.fromkeys(r["check"]), key=lambda c: int(c[1:])):
        dc = r[r["check"] == ck]
        body.append(label_row(CHECKS.get(ck, ck), 7))
        for v in dict.fromkeys(dc["variant"]):
            body += ame_rows(dc[dc["variant"] == v], prefix=f"{cap(v)}, ")
    caption(doc, "Table S7", "Robustness Checks")
    n = add_table(doc, AME_W, AME_HEAD, body, size=16)
    notes(doc, AME_NOTES + [
        "Each check repeats models of the main analysis with one change, named in the heading of the check. "
        "Check R4 (telephone interviews) concerns the waves 2020 and 2022, which the analysis does not use, "
        "and check R8 (years of schooling) is the decomposition that follows. Check R12 defines the stages "
        "S1 to S4 by the highest stage attended in place of the highest level completed. In check R1 and "
        "check R2, the decomposition of Table IV, panel A, is shown on the scale of the ordered logit "
        "coefficients. The indirect part as a percentage of the total is shown without a standard "
        "error, as in Table IV, because a ratio has no reliable standard error when the total is close "
        "to zero."] + count_note(r), size=9)
    o = rd("S7_oaxaca")
    widths = [3360, 1100, 1000, 1800, 1000, 1100]
    head = [[C(x) for x in ("", "Estimate", "S.E.", "95% C.I.", "p-value", "% of the gap")]]
    clab = {"total": "Gap, Han minus non-Han", "explained": "Explained by the terms",
            "unexplained": "Unexplained", "explained_agri3": "Hukou at age 3",
            "explained_cohort": "Birth cohort", "explained_male": "Gender",
            "explained_paredu": "Parents' education", "explained_province": "Province"}
    body = []
    for s in dict.fromkeys(o["sample"]):
        d = o[o["sample"] == s]
        r0 = d.iloc[0]
        body.append(label_row(f"{s}, years of schooling, n = {fmt_n(r0['n'])}, non-Han n = "
                              f"{fmt_n(r0['n_nonhan'])}, {r0['m']} imputed data sets", 6))
        for _, x in d.iterrows():
            body.append([C(clab.get(x["component"], x["component"]),
                           indent=2 * INDENT if x["component"].startswith("explained_") else INDENT),
                         C(fmt(x["estimate"], 3)), C(fmt(x["se"], 3)),
                         C(f"[{fmt(x['ci_low'], 3)}, {fmt(x['ci_high'], 3)}]"), C(fmt_p(x["p"])),
                         C(fmt(x["pct"], 1))])
    continued(doc, "Table S7", "Check R8, Decomposition of the Gap in Years of Schooling")
    n += add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Twofold decomposition by the method of Oaxaca and Blinder of the gap in years of schooling between "
        "Han and non-Han respondents, from weighted linear regressions with the terms of M4 in Table II "
        "without the interaction, with the coefficients of the Han respondents as the reference. The "
        "explained part is split by term. Years of schooling are in years. A province with fewer than 30 "
        "Han or 30 non-Han respondents joins the small-province group."], size=9)
    sel = rd("S7_r10_selection")
    widths = [1300, 900, 2100, 1000, 600, 900, 2560]
    head = [[C(x) for x in ("Sample", "Step", "Term", "Wald test", "df", "p-value", "Decision")]]
    body = [[C(x["sample"]), C(x["step"]), C(R10_TERMS.get(x["term"], cap(x["term"]))),
             C(fmt(x["wald"], 2)), C(x["df"] or DASH), C(fmt_p(x["p"])), C(cap(x["decision"]))]
            for _, x in sel.iterrows()]
    para(doc, [("Table S7 (continued)", True, False)], keep=True, before=12)
    para(doc, [("Check R10, Selection of the Terms Free across the Levels of Education", False, True)],
         keep=True, after=6)
    n += add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "A term made equal across levels has one coefficient for all levels of education. The terms of step "
        "0 were made equal across levels from the start, for the reason stated. At each later step, a Wald "
        "test compares, for each term still free, the model with the coefficients of the term free across "
        "the levels of education with the model with one coefficient for all levels. The term with the "
        "largest p-value of 0.05 or above is made equal across levels, and the selection stops when every "
        "term still free has a p-value below 0.05."], size=9)
    return n


def supp_imputation(doc):
    s8 = rd("S8_imputation")
    widths = [2160, 900, 900, 1100, 1000, 1100, 1100, 1100]
    head = [[C(x) for x in ("", "Group", "n", "Missing n", "Missing %", "Mean, observed",
                            "Mean, imputed", "Logged events")]]
    vl = {"paredu": "Parents' education (1 to 5)", "edu5": "Education (1 to 5)",
          "linc": "Log family income per person", "linc10": "Log family income per person, 2010"}
    body = []
    for s in dict.fromkeys(s8["sample"]):
        d = s8[s8["sample"] == s]
        body.append(label_row(f"{s}, {d.iloc[0]['m']} imputed data sets", 8))
        for _, x in d.iterrows():
            body.append([C(vl.get(x["variable"], x["variable"]), indent=INDENT), C(x["group"]),
                         C(fmt_n(x["n"])), C(fmt_n(x["n_missing_shown"])), C(fmt(x["pct_missing"], 1)),
                         C(fmt(x["mean_observed"], 2)), C(fmt(x["mean_imputed"], 2)),
                         C(x["logged_events"])])
    caption(doc, "Table S8", "Imputation of Missing Values")
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Missing values are imputed by chained equations (R package mice) within each sample, with 20 "
        "imputed data sets. The imputed mean is the mean of the imputed values over all data sets. Logged "
        "events is the number of events that mice recorded, for example a predictor left out of an "
        "imputation model because it was collinear with others.",
        "<10 marks a count from 1 to 9, suppressed marks a second count hidden so that the first cannot be "
        f"found by subtraction, and {DASH} marks the percentage and the imputed mean of such a count."],
        size=9)
    return n


def supp_attrition(doc):
    s9 = rd("S9_attrition")
    widths = [3960, 1500, 1300, 1200, 1400]
    head = [[C(x) for x in ("", "Status", "Estimate", "S.E.", "n")]]
    body = []
    for g in dict.fromkeys(s9["group"]):
        body.append(label_row(g, 5))
        for _, x in s9[s9["group"] == g].iterrows():
            body.append([C(x["variable"], indent=INDENT), C(x["status"]), C(fmt(x["estimate"], 2)),
                         C(fmt(x["se"], 2)), C(fmt_n(x["n_shown"]))])
    caption(doc, "Table S9", "Children Aged 12 to 15 in 2010, Followed Up to 2018 or Lost")
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Weighted percentages and means in 2010 of the children aged 12 to 15 in 2010, by whether they were "
        "interviewed again in 2018, with the 2010 cross-sectional weight and the survey design. n is "
        "unweighted.", SUPPRESS_NOTE], size=9)
    return n


def supp_classmates(doc):
    s10 = rd("S10_classmates")
    widths = [3160, 1500, 1500, 1500, 1700]
    head = [[C(x) for x in ("", "Group", "Estimate (%)", "S.E.", "n")]]
    il = {"ks3": "language of the classmates", "ks3m": "language of the classmates",
          "qs3m": "language of the classmates", "qz103": "main language of the interview"}
    body = []
    for w, it in dict.fromkeys(zip(s10["wave"], s10["item"])):
        d = s10[(s10["wave"] == w) & (s10["item"] == it)]
        body.append(label_row(f"{w}, {il.get(it, it)} (item {it})", 5))
        for _, x in d.iterrows():
            body.append([C(x["level"], indent=INDENT), C(x["group"]), C(fmt(x["estimate"], 1)),
                         C(fmt(x["se"], 1)), C(fmt_n(x["n_shown"]))])
    caption(doc, "Table S10", "Language of the Classmates and Main Language of the Interview")
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Weighted percentages with the cross-sectional weight of each wave and the survey design, among the "
        "respondents who answered the item. n is unweighted.", SUPPRESS_NOTE], size=9)
    return n


def s11_label(v):
    """Education: College or above (%) becomes Education, college or above (%)"""
    m = re.fullmatch(r"Education: (.*) \(%\)", v)
    return f"Education, {m.group(1)[0].lower() + m.group(1)[1:]} (%)" if m else v


def n_is(x):
    v = fmt_n(x)
    return {"<10": "n < 10", "suppressed": "n suppressed"}.get(v, f"n = {v}")


def supp_cohorts(doc):
    s11 = rd("S11_cohorts")
    widths = [2960, 1600, 1600, 1600, 1600]
    head = [[C(""), C("Han", 2), C("Non-Han", 2)],
            [C(""), C("CFPS 2010"), C("CFPS 2018"), C("CFPS 2010"), C("CFPS 2018")]]
    body = []
    for co in dict.fromkeys(s11["cohort"]):
        d = s11[s11["cohort"] == co]
        body.append(label_row(f"Born {co.replace('-', DASH)}", 5))
        for v in dict.fromkeys(d["variable"]):
            dec = 1 if "%" in v else 2
            cells = [C(s11_label(v), indent=INDENT)]
            for g in ("Han", "Non-Han"):
                for src in ("A-2010", "A-2018"):
                    x = ame_get(d, variable=v, group=g, source=src)
                    if x is None:
                        cells.append(C(DASH))
                    elif num(x["estimate"]) is None:
                        cells.append(C([DASH, n_is(x["n_shown"])]))
                    else:
                        cells.append(C([f"{fmt(x['estimate'], dec)} ({fmt(x['se'], dec)})",
                                        n_is(x["n_shown"])]))
            body.append(cells)
    caption(doc, "Table S11", "Education of the Same Birth Cohorts in CFPS 2010 and 2018")
    n = add_table(doc, widths, head, body, size=16)
    notes(doc, [
        "Weighted percentages and means, with standard errors in ( ), of the birth cohorts present in both "
        "samples A-2010 and A-2018, with the national cross-sectional weight of each wave and the survey "
        "design. n is unweighted.", SUPPRESS_NOTE], size=9)
    return n


#------------------------------------------------------------------------------
# Main
#------------------------------------------------------------------------------
def main():
    os.makedirs(OUT, exist_ok=True)
    LAST.update(last_cohorts())
    counts = {}

    doc = new_doc()
    counts["Table I"] = table_1(doc)
    counts["Table II"] = table_2(doc)
    counts["Table III"] = table_3(doc)
    counts["Table IV"] = table_4(doc)
    counts["Figure 1, points shown"] = figure_1()
    figure_page(doc)
    doc.save(os.path.join(OUT, "Main tables and figure.docx"))

    doc = new_doc()
    a = rd("S3_coef_A")
    bco = rd("S3_coef_B")
    counts["Table S1"] = supp_flow(doc)
    counts["Table S2"] = supp_groups(doc)
    counts["Table S3"] = supp_full(doc)
    counts["Table S4"] = supp_ame(
        doc, "Table S4", "Transitions T1 to T4, CFPS 2010 and 2018", rd("S4_transitions"),
        ["Logit models of each transition with the terms of M4 in Table II. T1 is completing primary school "
         "among all respondents, T2 junior high school among those who completed primary school, T3 senior "
         "high school among those who completed junior high school and T4 college or above among those who "
         "completed senior high school."],
        a[a["outcome"] != "edu5"])
    counts["Table S5"] = supp_ame(
        doc, "Table S5", "Education in CFPS 2012, 2014 and 2016", rd("S5_ame_other"),
        ["Models M1 to M4 of Table II and the transition models of Table S4, fitted to the samples of the "
         "waves 2012, 2014 and 2016."],
        rd("S3_coef_Aother"))
    counts["Table S6"] = supp_ame(
        doc, "Table S6", "School Enrolment at Ages 16 to 22, CFPS 2010 to 2018", rd("S6_ame_enrol"),
        ["Logit models of enrolment in school among the respondents aged 16 to 22 of the waves 2010 to 2018, "
         "pooled, with non-Han ethnicity, hukou at age 3 and their interaction, gender, parents' education, "
         "log family income per person, province, single years of age and the wave."],
        bco[bco["sample"].str.startswith("B-enrol")])
    counts["Table S7"] = supp_checks(doc)
    counts["Table S8"] = supp_imputation(doc)
    counts["Table S9"] = supp_attrition(doc)
    counts["Table S10"] = supp_classmates(doc)
    counts["Table S11"] = supp_cohorts(doc)
    doc.save(os.path.join(OUT, "Supplementary tables.docx"))

    for k, v in counts.items():
        print(f"{k}: {v} rows")
    print("Written to", OUT)


if __name__ == "__main__":
    main()
